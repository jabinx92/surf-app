from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import BaseDocTemplate, Frame, PageTemplate, Paragraph, Spacer


OUT_DIR = Path(r"C:\Users\jabin\surf-app\output\pdf")
PDF_OUTPUT = OUT_DIR / "Jonathan-Won-Full-Stack-Resume.pdf"
DOCX_OUTPUT = OUT_DIR / "Jonathan-Won-Full-Stack-Resume.docx"
OUT_DIR.mkdir(parents=True, exist_ok=True)

CONTACT_1 = "San Francisco, CA | 925-548-7768 | jabinx92@gmail.com"
CONTACT_2 = "jonathanwon.com | linkedin.com/in/jonathanwon | github.com/jabinx92"
SUMMARY = (
    "Full-Stack Software Engineer with 4.5 years of professional front-end development experience "
    "delivering responsive, production-facing web experiences for Gap brands across global markets. "
    "Builds end-to-end products with React, TypeScript, React Native, Expo, Supabase, PostgreSQL, SQL, "
    "REST APIs, authentication, database design, and CI/CD workflows."
)
SKILLS = [
    ("Languages & Frontend", "TypeScript, JavaScript (ES6+), React, React Native, Expo, HTML5, CSS3, SCSS, responsive web design"),
    ("Backend & Data", "Supabase, PostgreSQL, SQL, Node.js, Express.js, REST APIs, API integration, authentication, row-level security, Edge Functions"),
    ("Tools & Delivery", "Git, GitHub, GitHub Actions, CI/CD, EAS, Firebase, Optimizely, npm, Webpack, Chrome DevTools"),
]
EXPERIENCE = [
    ("Front-End Web Developer | Gap Inc. | May 2021 - Nov 2025 | San Francisco, CA", [
        "Delivered production-ready homepage and campaign experiences across mobile and desktop for Gap US, Canada, Factory, Japan, and Europe.",
        "Built and maintained reusable front-end solutions with React, JavaScript, HTML, CSS, SCSS, JSON, and jQuery in a large-scale retail environment.",
        "Partnered with designers, product managers, marketers, and cross-functional teams to translate business requirements into responsive, accessible customer experiences.",
        "Implemented A/B testing with Optimizely and supported debugging, quality assurance, and production releases across regional storefronts.",
    ])
]
PROJECTS = [
    ("Founder & Full-Stack Developer | WaveLens | 2026 - Present", ("wavelens.expo.app", "github.com/jabinx92/surf-app"), [
        "Built a cross-platform marketplace with Expo, React Native, and TypeScript that connects surfers with local photographers and videographers.",
        "Developed authentication, creator applications and approvals, editable profiles and portfolios, availability, booking requests, messaging, notifications, reviews, reporting, and admin moderation.",
        "Designed the Supabase and PostgreSQL backend with SQL migrations, row-level security, storage, and an account-deletion Edge Function; configured EAS environments and GitHub Actions CI/CD automation.",
    ]),
    ("Developer | JonathanWon.com | Personal Portfolio", None, [
        "Built and deployed a responsive Gatsby and React portfolio showcasing professional experience and selected software projects, with a custom domain and Git-based deployment workflow.",
    ]),
]


def build_pdf():
    navy = colors.HexColor("#102A43")
    blue = colors.HexColor("#147D92")
    slate = colors.HexColor("#486581")
    text = colors.HexColor("#243B53")
    styles = getSampleStyleSheet()
    name = ParagraphStyle("Name", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=20, leading=21, textColor=navy, alignment=TA_CENTER, spaceAfter=2)
    target = ParagraphStyle("Target", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=9.5, leading=11, textColor=blue, alignment=TA_CENTER, spaceAfter=2)
    contact = ParagraphStyle("Contact", parent=styles["Normal"], fontName="Helvetica", fontSize=8.3, leading=10.3, textColor=slate, alignment=TA_CENTER)
    section = ParagraphStyle("Section", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=10.5, leading=12, textColor=blue, spaceBefore=6.5, spaceAfter=2.5, borderWidth=0, borderPadding=0)
    body = ParagraphStyle("Body", parent=styles["Normal"], fontName="Helvetica", fontSize=8.85, leading=11.25, textColor=text)
    role = ParagraphStyle("Role", parent=body, fontName="Helvetica-Bold", fontSize=9.2, leading=11.4, textColor=navy, spaceBefore=1, spaceAfter=1)
    bullet = ParagraphStyle("Bullet", parent=body, leftIndent=11, firstLineIndent=-7, spaceBefore=1)
    skill = ParagraphStyle("Skill", parent=body, fontSize=8.65, leading=10.9)
    link = ParagraphStyle("Link", parent=contact, alignment=0, spaceAfter=1)

    def heading(value):
        return Paragraph(value, section)

    def item(value):
        # ASCII marker keeps PDF extraction clean for older ATS parsers.
        return Paragraph(f"- {value}", bullet)

    doc = BaseDocTemplate(str(PDF_OUTPUT), pagesize=letter, leftMargin=.58*inch, rightMargin=.58*inch, topMargin=.35*inch, bottomMargin=.32*inch, title="Jonathan Won - Full-Stack Software Engineer Resume", author="Jonathan Won")
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    doc.addPageTemplates(PageTemplate(id="resume", frames=[frame]))
    story = [
        Paragraph("Jonathan Won", name),
        Paragraph("FULL-STACK SOFTWARE ENGINEER", target),
        Paragraph(CONTACT_1, contact),
        Paragraph(CONTACT_2, contact),
        Spacer(1, 3),
        heading("PROFESSIONAL SUMMARY"), Paragraph(SUMMARY, body),
        heading("TECHNICAL SKILLS"),
    ]
    for label, value in SKILLS:
        story.append(Paragraph(f"<b>{label}:</b> {value}", skill))
    story.append(heading("WORK EXPERIENCE"))
    for label, bullets in EXPERIENCE:
        story.append(Paragraph(label, role))
        story.extend(item(v) for v in bullets)
    story.append(heading("PROJECTS"))
    for label, urls, bullets in PROJECTS:
        story.append(Paragraph(label, role))
        if urls:
            live_url, github_url = urls
            story.append(Paragraph(f"<b>Live demo:</b> <link href='https://{live_url}' color='#147D92'>{live_url}</link> &nbsp;|&nbsp; <b>Source code:</b> <link href='https://{github_url}' color='#147D92'>{github_url}</link>", link))
        story.extend(item(v) for v in bullets)
    story.extend([
        heading("ADDITIONAL EXPERIENCE"),
        Paragraph("<b>Order Processor Specialist, Global Syn-Turf</b> | 2018 - 2019 | <b>Operations Coordinator, Comba Telecom</b> | 2017 - 2018", body),
        heading("EDUCATION"),
        Paragraph("<b>San Francisco State University</b> | Bachelor of Science in Business Marketing", body),
    ])
    doc.build(story)


def set_font(run, size, bold=False, color="000000"):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_bottom_border(paragraph, color="B8C4CE"):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def add_hyperlink(paragraph, text, url):
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "147D92")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend([color, underline])
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.extend([run_properties, text_element])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(.38)
    section.bottom_margin = Inches(.35)
    section.left_margin = Inches(.58)
    section.right_margin = Inches(.58)
    section.header_distance = Inches(.3)
    section.footer_distance = Inches(.3)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(8.9)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.color.rgb = RGBColor(20, 125, 146)
    h1 = doc.styles["Heading 1"]
    h1.font.size = Pt(10.5)
    h1.font.bold = True
    h1.paragraph_format.space_before = Pt(6)
    h1.paragraph_format.space_after = Pt(2)
    h1.paragraph_format.keep_with_next = True

    bullet_style = doc.styles["List Bullet"]
    bullet_style.font.name = "Arial"
    bullet_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    bullet_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    bullet_style.font.size = Pt(8.8)
    bullet_style.paragraph_format.left_indent = Inches(.18)
    bullet_style.paragraph_format.first_line_indent = Inches(-.13)
    bullet_style.paragraph_format.space_after = Pt(1)
    bullet_style.paragraph_format.line_spacing = 1.0

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run("Jonathan Won"), 20, True, "102A43")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    set_font(p.add_run("FULL-STACK SOFTWARE ENGINEER"), 9.5, True, "147D92")
    for value in (CONTACT_1, CONTACT_2):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(value), 8.3, False, "486581")

    def heading(value):
        p = doc.add_paragraph(value, style="Heading 1")
        add_bottom_border(p)

    def paragraph(value, size=8.9):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(value), size)

    def labeled(label, value):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(label + ": "), 8.7, True)
        set_font(p.add_run(value), 8.7)

    def role(value):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.keep_with_next = True
        set_font(p.add_run(value), 9.2, True, "102A43")

    def bullet(value):
        p = doc.add_paragraph(style="List Bullet")
        set_font(p.add_run(value), 8.8)

    heading("PROFESSIONAL SUMMARY")
    paragraph(SUMMARY)
    heading("TECHNICAL SKILLS")
    for label, value in SKILLS:
        labeled(label, value)
    heading("WORK EXPERIENCE")
    for label, bullets in EXPERIENCE:
        role(label)
        for value in bullets:
            bullet(value)
    heading("PROJECTS")
    for label, urls, bullets in PROJECTS:
        role(label)
        if urls:
            live_url, github_url = urls
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            set_font(p.add_run("Live demo: "), 8.2, True)
            add_hyperlink(p, live_url, f"https://{live_url}")
            set_font(p.add_run(" | Source code: "), 8.2, True)
            add_hyperlink(p, github_url, f"https://{github_url}")
        for value in bullets:
            bullet(value)
    heading("ADDITIONAL EXPERIENCE")
    paragraph("Order Processor Specialist, Global Syn-Turf | 2018 - 2019 | Operations Coordinator, Comba Telecom | 2017 - 2018", 8.6)
    heading("EDUCATION")
    paragraph("San Francisco State University | Bachelor of Science in Business Marketing", 8.7)

    props = doc.core_properties
    props.title = "Jonathan Won - Full-Stack Software Engineer Resume"
    props.author = "Jonathan Won"
    props.subject = "Full-Stack Software Engineer resume"
    props.keywords = "Full-Stack Software Engineer, React, TypeScript, React Native, Expo, Supabase, PostgreSQL, SQL, REST APIs, CI/CD"
    doc.save(DOCX_OUTPUT)


build_pdf()
build_docx()
print(PDF_OUTPUT)
print(DOCX_OUTPUT)
